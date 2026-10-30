"""Сборка руководств «Пульса Проекта» из Markdown в DOCX и PDF.

    python3 tools/docs/build_manuals.py              # все руководства
    python3 tools/docs/build_manuals.py user-guide   # одно

Нужны pandoc 3.x, python-docx, pypdf, LibreOffice (оглавление DOCX) и Node.js с playwright (Chromium).
Источник — docs/manuals/*.md и docs/manuals/images; результат кладётся рядом с .md.
PDF собирается в два прохода: после первого номера страниц разделов читаются из ссылок
содержания и вписываются в него во втором. Оформление и порядок сборки — как у todo-planner.
"""
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[2]
MANUALS = ROOT / "docs" / "manuals"
TOOLS = Path(__file__).resolve().parent
VERSION = json.loads((ROOT / "package.json").read_text(encoding="utf-8"))["version"]
DOCS = {
    "user-guide": ("Руководство пользователя", f"Пульс Проекта {VERSION} · Руководство пользователя"),
    "admin-guide": ("Руководство администратора", f"Пульс Проекта {VERSION} · Руководство администратора"),
    "deployment-guide": ("Руководство по развёртыванию", f"Пульс Проекта {VERSION} · Руководство по развёртыванию"),
}
BLUE = RGBColor(0x2F, 0x4F, 0xD1)
INK = RGBColor(0x1D, 0x24, 0x33)
FONT = "Arial"


def run(cmd, **kw):
    subprocess.run(cmd, check=True, **kw)


def source_dir(_name):
    return MANUALS


# ---------------------------------------------------------------- PDF
def toc_page_numbers(pdf_path):
    """Номера страниц, на которые ведут ссылки содержания (в порядке ссылок)."""
    reader = PdfReader(str(pdf_path))
    page_index = {page.indirect_reference.idnum: i for i, page in enumerate(reader.pages)}
    named = {}
    try:
        for key, dest in reader.named_destinations.items():
            named[key] = reader.get_destination_page_number(dest)
    except Exception:  # noqa: BLE001 — разные версии Chromium пишут назначения по-разному
        pass
    numbers = []
    for page in reader.pages[1:6]:  # содержание сразу после обложки
        annots = page.get("/Annots") or []
        links = []
        for ref in annots:
            annot = ref.get_object()
            if annot.get("/Subtype") != "/Link":
                continue
            dest = annot.get("/Dest")
            if dest is None and annot.get("/A"):
                dest = annot["/A"].get("/D")
            target = None
            if isinstance(dest, str) or hasattr(dest, "startswith"):
                target = named.get(str(dest))
            elif dest is not None:
                target = page_index.get(dest[0].idnum)
            if target is None:
                continue
            rect = annot.get("/Rect")
            links.append((-float(rect[3]), float(rect[0]), target + 1))
        if not links and numbers:
            break
        numbers.extend(n for *_pos, n in sorted(links))
    return numbers


def replace_cover(pdf, cover):
    """Первая страница — обложка без колонтитула; закладки и ссылки остальных страниц сохраняются."""
    writer = PdfWriter(clone_from=str(pdf))
    writer.insert_page(PdfReader(str(cover)).pages[0], 0)
    writer.remove_page(1)
    writer.write(str(pdf))


def build_pdf(name, md, title, footer):
    html = md.with_suffix(".build.html")
    css = (TOOLS / "manual.css").read_text(encoding="utf-8")
    run(["pandoc", str(md), "-f", "markdown+implicit_figures+link_attributes+ascii_identifiers", "-t", "html5",
         "--template", str(TOOLS / "template.html"), "--toc", "--toc-depth=2", "--number-sections",
         "-V", f"css={css}", "-V", f"cover-title={title}", "--resource-path", str(md.parent), "-o", str(html)])
    pdf = md.with_suffix(".pdf")
    script = str(TOOLS / "html_to_pdf.mjs")
    try:
        run(["node", script, str(html), str(pdf), footer], cwd=ROOT)
        numbers = toc_page_numbers(pdf)
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False) as fh:
            json.dump(numbers, fh)
        run(["node", script, str(html), str(pdf), footer, fh.name], cwd=ROOT)
        Path(fh.name).unlink()
        cover = pdf.with_suffix(".cover.pdf")
        run(["node", script, str(html), str(cover), footer], cwd=ROOT, env={**os.environ, "COVER_ONLY": "1"})
        replace_cover(pdf, cover)
        cover.unlink()
    finally:
        html.unlink(missing_ok=True)
    print(f"  pdf: {pdf.relative_to(ROOT)} ({len(PdfReader(str(pdf)).pages)} стр., ссылок содержания: {len(numbers)})")


# ---------------------------------------------------------------- DOCX
def reference_docx(path):
    data = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], check=True, capture_output=True).stdout
    path.write_bytes(data)
    doc = Document(str(path))
    by_name = {s.name: s for s in doc.styles}

    class _Styles:  # pandoc называет стили «Heading 1», python-docx ищет «heading 1»
        def __getitem__(self, key):
            return by_name[key]

        def __iter__(self):
            return iter(by_name.values())

    styles = _Styles()

    def font(style, size=None, bold=None, color=None, name=FONT):
        style.font.name = name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), name)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            rfonts.attrib.pop(qn(attr), None)
        if size:
            style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if color is not None:
            style.font.color.rgb = color

    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in styles]:
            font(styles[name], 10.5, color=INK)
            styles[name].paragraph_format.space_after = Pt(5)
            styles[name].paragraph_format.line_spacing = 1.15
    font(styles["Title"], 28, True, INK)
    font(styles["Subtitle"], 15, False, BLUE)
    if "Date" in [s.name for s in styles]:
        font(styles["Date"], 11, False, RGBColor(0x6B, 0x73, 0x85))
    sizes = {1: 18, 2: 14, 3: 12, 4: 11}
    for level, size in sizes.items():
        style = styles[f"Heading {level}"]
        font(style, size, True, BLUE if level == 2 else INK)
        style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = True
    for name in ("Image Caption", "Caption", "Table Caption"):
        if name in [s.name for s in styles]:
            font(styles[name], 9, False, RGBColor(0x6B, 0x73, 0x85))
            styles[name].font.italic = False
    if "Source Code" in [s.name for s in styles]:
        font(styles["Source Code"], 8.5, name="Consolas")
    if "Block Text" in [s.name for s in styles]:
        font(styles["Block Text"], 10, False, INK)
    if "TOC Heading" in [s.name for s in styles]:
        font(styles["TOC Heading"], 18, True, INK)
    doc.save(str(path))


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table):
    tbl_pr = table._tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "D6DCEA")
        el.append(b)
    tbl_pr.append(el)


def polish_docx(path, product):
    doc = Document(str(path))
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        borders(table)
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                if i == 0:
                    shade(cell, "EEF2FB")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    for r in paragraph.runs:
                        r.font.size = Pt(9)
                        if i == 0:
                            r.font.bold = True
    # Выноски (цитаты) — голубая заливка
    for paragraph in doc.paragraphs:
        if paragraph.style.name in ("Block Text",):
            p_pr = paragraph._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F2F5FF")
            p_pr.append(shd)
    # Колонтитул с номером страницы
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.text = f"{product} · стр. "
    run = footer.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            run._r.append(instr)
    for r in footer.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x8A, 0x92, 0xA6)
    doc.save(str(path))


def build_docx(name, md, title, footer, reference):
    out = md.with_suffix(".docx")
    run(["pandoc", str(md), "-f", "markdown+implicit_figures+link_attributes", "-t", "docx",
         "--reference-doc", str(reference), "--toc", "--toc-depth=2", "--number-sections",
         "-M", "toc-title=Содержание",
         "--resource-path", str(md.parent), "-o", str(out)])
    # Колонтитул DOCX — та же версия, что в колонтитуле PDF
    polish_docx(out, footer.split(" · ")[0])
    # Оглавление с номерами страниц: LibreOffice обновляет поле TOC и сохраняет DOCX
    run([sys.executable, str(TOOLS / "update_docx_toc.py"), str(out)])
    print(f"  docx: {out.relative_to(ROOT)}")


def main(names):
    with tempfile.TemporaryDirectory() as tmp:
        reference = Path(tmp) / "reference.docx"
        reference_docx(reference)
        for name in names:
            title, footer = DOCS[name]
            md = source_dir(name) / f"{name}.md"
            print(name)
            build_docx(name, md, title, footer, reference)
            build_pdf(name, md, title, footer)


if __name__ == "__main__":
    main(sys.argv[1:] or [n for n in DOCS if (source_dir(n) / f"{n}.md").exists()])
